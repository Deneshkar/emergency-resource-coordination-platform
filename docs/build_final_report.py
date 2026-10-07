import docx
import os
import sys

# Ensure local imports work
sys.path.append(os.path.dirname(os.path.abspath(__file__)))
from part1_builder import create_full_part1_elements

def generate_final_report():
    source_path = 'docs/SE3090_Project_Report_SEF_backup.docx'
    final_output_path = 'docs/SE3090_Project_Report_SEF.docx'

    print(f"Loading base document from: {source_path}")
    doc = docx.Document(source_path)
    body = doc._body._element
    body_children = body.xpath('./w:p | ./w:tbl')

    # Locate the boundary where Part II starts
    part2_idx = None
    for i, el in enumerate(body_children):
        if el.tag.endswith('p'):
            text = "".join(el.itertext())
            if "PART II: INDIVIDUAL REPORTS" in text:
                part2_idx = i
                break

    if part2_idx is None:
        raise ValueError("Could not find 'PART II: INDIVIDUAL REPORTS' anchor in document!")

    print(f"Located 'PART II: INDIVIDUAL REPORTS' at child index {part2_idx} out of {len(body_children)}.")
    old_part1_elements = body_children[:part2_idx]
    part2_anchor = body_children[part2_idx]

    # Track existing elements
    initial_elements = set(body.xpath('./w:p | ./w:tbl'))

    # Generate all new Part I elements with complete sections, tables, and images
    print("Generating comprehensive Part I Group Technical & Architectural Report...")
    create_full_part1_elements(doc)

    # Collect newly added elements
    current_elements = body.xpath('./w:p | ./w:tbl')
    new_elements = [el for el in current_elements if el not in initial_elements]
    print(f"Generated {len(new_elements)} new elements for Part I.")

    # Move new elements before Part II anchor
    print("Re-ordering elements so new Part I precedes Part II...")
    for el in new_elements:
        body.remove(el)
        part2_anchor.addprevious(el)

    # Remove old Part I elements
    print(f"Removing {len(old_part1_elements)} obsolete Part I elements...")
    for el in old_part1_elements:
        body.remove(el)

    # ==================== PART II REFINEMENTS ====================
    print("Applying individual declaration refinements and date corrections...")
    
    # 1. Student 2 Declaration & Signature
    for p in doc.paragraphs:
        if "I declare that the work described in this Individual Report section is my own contribution" in p.text:
            p.text = (
                "I, M. Thanusikan (IT24103609), declare that the work presented in my individual contribution to "
                "the SE3090 Assignment 1 is my own work and that I have contributed directly to the architecture, "
                "development, testing, and deployment of the Emergency Resource Coordination Platform as detailed herein. "
                "All code, database schemas, test cases, and documentation have been reviewed, executed, and verified by me. "
                "I can explain, modify, and defend all submitted components during the viva evaluation."
            )
            print("  [Updated] Student 2 declaration text paragraph.")
            
        if "Student Identifier: Student 2" in p.text and "Assigned Role:" in p.text:
            # We preserve run 5 (which contains the signature drawing)
            if len(p.runs) >= 6:
                p.runs[0].text = "Student Name: M. Thanusikan"
                p.runs[1].text = "\nStudent ID: IT24103609\nAssigned Role: Student 2 — Resource & Inventory Lead + Situation Analysis Agent"
                p.runs[2].text = "\nDate: 28 September 2026"
                p.runs[3].text = "\n"
                p.runs[4].text = "\nSignature: "
                # run 5 is the signature drawing
                p.runs[6].text = " (M. Thanusikan)"
                print("  [Updated] Student 2 signature block while preserving drawing in Run 5.")
            else:
                p.text = (
                    "Student Name: M. Thanusikan\n"
                    "Student ID: IT24103609\n"
                    "Assigned Role: Student 2 — Resource & Inventory Lead + Situation Analysis Agent\n"
                    "Date: 28 September 2026\n\n"
                    "Signature: M. Thanusikan"
                )
                print("  [Updated] Student 2 signature block text.")

    # 2. Student 3 Date Inconsistency Fix (01.10.2026 -> 30.09.2026)
    for p in doc.paragraphs:
        if "01.10.2026" in p.text:
            for r in p.runs:
                if "01.10.2026" in r.text:
                    r.text = r.text.replace("01.10.2026", "30.09.2026")
            if "01.10.2026" in p.text:
                p.text = p.text.replace("01.10.2026", "30.09.2026")
            print("  [Updated] Student 3 declaration date from 01.10.2026 to 30.09.2026.")

    # Save to the authoritative file
    print(f"Saving final consolidated document to: {final_output_path}")
    doc.save(final_output_path)
    print("Done! Report successfully generated and saved.")

if __name__ == '__main__':
    generate_final_report()
