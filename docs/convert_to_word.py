import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import re

def set_cell_background(cell, fill_color):
    """Set cell shading color (hex string like 'EBF8FF')"""
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set internal cell padding in dxa (1 pt = 20 dxa)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_report_docx():
    doc = docx.Document()

    # Set page margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.8)
        s.bottom_margin = Inches(0.8)
        s.left_margin = Inches(0.8)
        s.right_margin = Inches(0.8)

    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    # Read Markdown source
    with open('docs/Consolidated_Assignment_Report.md', 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_table = False
    table_lines = []
    in_code_block = False
    code_block_lines = []

    def flush_table(t_lines):
        if not t_lines:
            return
        # Parse table lines
        rows = []
        for l in t_lines:
            l = l.strip()
            if not l or l.startswith('| :---') or l.startswith('|:---') or l.startswith('|---'):
                continue
            cells = [c.strip() for c in l.split('|')[1:-1]]
            if cells:
                rows.append(cells)
        if not rows:
            return

        num_cols = max(len(r) for r in rows)
        tbl = doc.add_table(rows=len(rows), cols=num_cols)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl.autofit = True

        for r_idx, row_data in enumerate(rows):
            is_header = (r_idx == 0)
            row = tbl.rows[r_idx]
            for c_idx in range(num_cols):
                cell_text = row_data[c_idx] if c_idx < len(row_data) else ""
                cell = row.cells[c_idx]
                cell.text = cell_text
                set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
                
                # Format cell text
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                if p.runs:
                    run = p.runs[0]
                    run.font.name = 'Calibri'
                    run.font.size = Pt(9.5) if is_header else Pt(9)
                    if is_header:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
                
                if is_header:
                    set_cell_background(cell, "EBF8FF")
                elif r_idx % 2 == 1:
                    set_cell_background(cell, "F7FAFC")
                else:
                    set_cell_background(cell, "FFFFFF")
        
        # Add space after table
        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_before = Pt(4)
        p_after.paragraph_format.space_after = Pt(6)

    def flush_code_block(c_lines):
        if not c_lines:
            return
        code_text = "".join(c_lines).strip()
        # Add code container box
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_background(cell, "1A202C")
        set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run(code_text)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0xFA, 0xFA, 0xFA)

        p_after = doc.add_paragraph()
        p_after.paragraph_format.space_after = Pt(6)

    for line in lines:
        raw_line = line
        line = line.strip()

        # Handle page breaks
        if '<div style="page-break-after: always;"></div>' in line or line == '\\pagebreak':
            if in_table:
                flush_table(table_lines)
                table_lines = []
                in_table = False
            doc.add_page_break()
            continue

        # Handle code blocks
        if line.startswith('```'):
            if in_code_block:
                flush_code_block(code_block_lines)
                code_block_lines = []
                in_code_block = False
            else:
                if in_table:
                    flush_table(table_lines)
                    table_lines = []
                    in_table = False
                in_code_block = True
            continue

        if in_code_block:
            code_block_lines.append(raw_line)
            continue

        # Handle tables
        if line.startswith('|'):
            in_table = True
            table_lines.append(line)
            continue
        else:
            if in_table:
                flush_table(table_lines)
                table_lines = []
                in_table = False

        # Handle screenshot placeholders
        if line.startswith('[SCREENSHOT:') and line.endswith(']'):
            desc = line[12:-1].strip()
            # Add styled callout box for screenshot
            tbl = doc.add_table(rows=1, cols=1)
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell = tbl.cell(0, 0)
            set_cell_background(cell, "F0F9FF")
            set_cell_margins(cell, top=200, bottom=200, left=240, right=240)
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(2)
            run_title = p.add_run("📷  [PASTE SCREENSHOT HERE]\n")
            run_title.font.name = 'Calibri'
            run_title.font.size = Pt(11)
            run_title.font.bold = True
            run_title.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

            run_desc = p.add_run(desc)
            run_desc.font.name = 'Calibri'
            run_desc.font.size = Pt(10)
            run_desc.font.italic = True
            run_desc.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

            p_after = doc.add_paragraph()
            p_after.paragraph_format.space_after = Pt(6)
            continue

        # Handle headings
        if line.startswith('# '):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(18)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(line[2:])
            run.font.name = 'Calibri'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1A, 0x36, 0x5D)
            continue

        if line.startswith('## '):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(line[3:])
            run.font.name = 'Calibri'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
            continue

        if line.startswith('### '):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(line[4:])
            run.font.name = 'Calibri'
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
            continue

        if line.startswith('#### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(line[5:])
            run.font.name = 'Calibri'
            run.font.size = Pt(11)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
            continue

        # Handle horizontal divider
        if line in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run("―" * 50)
            run.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE0)
            continue

        # Handle blank lines
        if not line:
            continue

        # Handle bullet lists
        if line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            # parse bold/code
            add_formatted_text(p, line[2:])
            continue

        # Handle numbered lists
        num_match = re.match(r'^(\d+)\.\s+(.*)', line)
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)
            add_formatted_text(p, num_match.group(2))
            continue

        # Regular paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(4)
        add_formatted_text(p, line)

    if in_table:
        flush_table(table_lines)
    if in_code_block:
        flush_code_block(code_block_lines)

    output_path = 'docs/Consolidated_Assignment_Report.docx'
    doc.save(output_path)
    print(f"Successfully generated {output_path}!")

def add_formatted_text(paragraph, text):
    """Parses simple inline markdown (bold, italic, code) into Word runs"""
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for tok in tokens:
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**') and len(tok) >= 4:
            run = paragraph.add_run(tok[2:-2])
            run.bold = True
            run.font.name = 'Calibri'
        elif tok.startswith('*') and tok.endswith('*') and len(tok) >= 2:
            run = paragraph.add_run(tok[1:-1])
            run.italic = True
            run.font.name = 'Calibri'
        elif tok.startswith('`') and tok.endswith('`') and len(tok) >= 2:
            run = paragraph.add_run(tok[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xC5, 0x30, 0x30)
        else:
            run = paragraph.add_run(tok)
            run.font.name = 'Calibri'

if __name__ == '__main__':
    create_report_docx()
