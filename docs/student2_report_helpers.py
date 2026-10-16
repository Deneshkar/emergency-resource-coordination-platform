import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    """Set the background color of a table cell."""
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set inner margins (padding) for a table cell in twips."""
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def add_callout(doc, title, text, bg_hex="F0F4F8", border_hex="1B365D"):
    """Add a professional styled callout box."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, bg_hex)
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

    # Set left border thick, others none
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="{border_hex}"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(4)
    run_t = p.add_run(f"📌 {title}\n")
    run_t.bold = True
    run_t.font.name = "Calibri"
    run_t.font.size = Pt(11)
    run_t.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    run_b = p.add_run(text)
    run_b.font.name = "Calibri"
    run_b.font.size = Pt(10)
    run_b.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_str, language="C# / Python"):
    """Add a code block with light gray background and monospace font."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.5)
    set_cell_background(cell, "F8F9FA")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)

    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="CBD5E0"/>'
        f'<w:left w:val="single" w:sz="18" w:space="0" w:color="4A5568"/>'
        f'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="CBD5E0"/>'
        f'<w:right w:val="single" w:sz="6" w:space="0" w:color="CBD5E0"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)

    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05

    run_lang = p.add_run(f"// {language} Snippet\n")
    run_lang.font.name = "Consolas"
    run_lang.font.size = Pt(9)
    run_lang.font.color.rgb = RGBColor(0x71, 0x80, 0x96)
    run_lang.italic = True

    run_code = p.add_run(code_str)
    run_code.font.name = "Consolas"
    run_code.font.size = Pt(9)
    run_code.font.color.rgb = RGBColor(0x1A, 0x20, 0x2C)

    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(4)

def add_figure(doc, img_path, figure_num, caption, desc, width_inches=6.0):
    """Add a real screenshot image with caption and technical description."""
    if not os.path.exists(img_path):
        print(f"Warning: Image path not found: {img_path}")
        return

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width_inches))

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(2)
    run_fig = p_cap.add_run(f"Figure {figure_num}: ")
    run_fig.bold = True
    run_fig.font.name = "Calibri"
    run_fig.font.size = Pt(9.5)
    run_fig.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    run_cap = p_cap.add_run(caption)
    run_cap.font.name = "Calibri"
    run_cap.font.size = Pt(9.5)
    run_cap.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    p_desc = doc.add_paragraph()
    p_desc.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_desc.paragraph_format.space_before = Pt(0)
    p_desc.paragraph_format.space_after = Pt(8)
    run_desc = p_desc.add_run(desc)
    run_desc.italic = True
    run_desc.font.name = "Calibri"
    run_desc.font.size = Pt(9)
    run_desc.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)

def format_table_header(row, col_names, bg_hex="1B365D"):
    """Format table header row with custom color and bold white text."""
    for i, name in enumerate(col_names):
        cell = row.cells[i]
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=100, bottom=100, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(name)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

def populate_table_row(row, values, is_even=False):
    """Populate a table row with zebra shading and clean text."""
    bg_hex = "F7FAFC" if is_even else "FFFFFF"
    for i, val in enumerate(values):
        cell = row.cells[i]
        set_cell_background(cell, bg_hex)
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.05
        run = p.add_run(val)
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

print("Builder helper functions declared successfully.")
