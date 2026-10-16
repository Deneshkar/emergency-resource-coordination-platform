import docx

doc = docx.Document('docs/SE3090_Project_Report_SEF.docx')
print("=== VERIFICATION REPORT ===")
print("Paragraphs:", len(doc.paragraphs))
print("Tables:", len(doc.tables))
print("Inline shapes:", len(doc.inline_shapes))

print("\n--- Declarations & Signatures ---")
for i, p in enumerate(doc.paragraphs):
    if any(k in p.text.lower() for k in ['declaration', 'date:', 'signature:']):
        print(f"P{i}: {p.text.strip()[:160]}")

print("\n--- Captions (Figures & Tables) ---")
for p in doc.paragraphs:
    if p.text.startswith(('Figure ', 'Table ')):
        print(p.text.strip())

# Check tables in cells for captions (side-by-side mobile figures)
for t in doc.tables:
    for r in t.rows:
        for c in r.cells:
            for p in c.paragraphs:
                if p.text.startswith('Figure '):
                    print("  [Table Cell]", p.text.strip())
